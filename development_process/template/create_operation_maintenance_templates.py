from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

OUTPUT_DIR = Path(__file__).parent

NAVY = "1F4E78"
BLUE = "D9EAF7"
LIGHT = "F3F6F9"
WHITE = "FFFFFF"
GRAY = "666666"
THIN = Side(style="thin", color="B7C9D6")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def style_sheet(ws, title, purpose, headers, rows, widths=None, status_col=None):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    ws.cell(1, 1, title)
    ws.cell(1, 1).font = Font(bold=True, size=16, color=WHITE)
    ws.cell(1, 1).fill = PatternFill("solid", fgColor=NAVY)
    ws.cell(1, 1).alignment = Alignment(horizontal="left")

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
    ws.cell(2, 1, purpose)
    ws.cell(2, 1).font = Font(italic=True, color=GRAY)
    ws.cell(2, 1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[2].height = 32

    metadata = ["案件名", "システム名", "環境", "作成日", "版", "作成者", "承認者"]
    for col, label in enumerate(metadata, 1):
        ws.cell(3, col, label)
        ws.cell(3, col).font = Font(bold=True, color=NAVY)
        ws.cell(3, col).fill = PatternFill("solid", fgColor=BLUE)
        ws.cell(3, col).border = BORDER
        if col <= len(headers):
            ws.cell(4, col, "")
            ws.cell(4, col).border = BORDER
    header_row = 6
    for col, header in enumerate(headers, 1):
        cell = ws.cell(header_row, col, header)
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = BORDER
    ws.row_dimensions[header_row].height = 32

    for r, row in enumerate(rows, header_row + 1):
        for c, value in enumerate(row, 1):
            cell = ws.cell(r, c, value)
            cell.border = BORDER
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if r % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=LIGHT)
    if rows:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(headers))}{header_row + len(rows)}"
    ws.freeze_panes = "A7"
    ws.sheet_view.showGridLines = False
    for col, width in enumerate(widths or [18] * len(headers), 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    if status_col:
        dv = DataValidation(type="list", formula1='"未着手,対応中,確認待ち,完了,保留"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{get_column_letter(status_col)}7:{get_column_letter(status_col)}200")


def add_readme(wb, name, purpose, sheets):
    ws = wb.create_sheet("使い方", 0)
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 95
    ws["A1"] = name
    ws["A1"].font = Font(bold=True, size=16, color=WHITE)
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws.merge_cells("A1:B1")
    ws["A3"] = "用途"
    ws["B3"] = purpose
    ws["A4"] = "入力方法"
    ws["B4"] = "青色の見出しに対応する行へ入力してください。例示行は自案件の内容に置き換えて使用します。"
    ws["A5"] = "必須確認"
    ws["B5"] = "案件名、環境、版、作成者、承認者、更新日、状態、判断・証跡を記録してください。"
    ws["A7"] = "シート"
    ws["B7"] = "内容"
    for cell in ws[7]:
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.border = BORDER
    for row, (sheet, desc) in enumerate(sheets, 8):
        ws.cell(row, 1, sheet)
        ws.cell(row, 2, desc)
        ws.cell(row, 1).border = BORDER
        ws.cell(row, 2).border = BORDER
        ws.cell(row, 2).alignment = Alignment(wrap_text=True)
    for row in range(3, 6):
        ws.cell(row, 1).font = Font(bold=True, color=NAVY)
        ws.cell(row, 1).fill = PatternFill("solid", fgColor=BLUE)
        ws.cell(row, 1).border = BORDER
        ws.cell(row, 2).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row, 2).border = BORDER
    ws.freeze_panes = "A8"


def save_workbook(filename, name, purpose, sheet_specs):
    wb = Workbook()
    wb.remove(wb.active)
    add_readme(wb, name, purpose, [(spec[0], spec[1]) for spec in sheet_specs])
    for sheet_name, sheet_purpose, headers, rows, widths, status_col in sheet_specs:
        ws = wb.create_sheet(sheet_name)
        style_sheet(ws, sheet_name, sheet_purpose, headers, rows, widths, status_col)
    wb.save(OUTPUT_DIR / filename)


save_workbook(
    "01_operation_runbook_template.xlsx",
    "運用設計書（Runbook）",
    "定常運用、障害一次対応、復旧手順を、担当者が同じ手順で実施できるように定義する。",
    [
        ("手順一覧", "運用手順の全体を一覧管理する。", ["手順ID", "手順名", "分類", "対象システム/機能", "実施タイミング", "実施者", "前提・権限", "完了条件", "関連手順", "最終更新日", "状態"], [["RUN-001", "日次バックアップ確認", "定常", "業務DB", "毎営業日", "運用担当", "運用アカウント", "バックアップ成功を確認", "RUN-002", "", "未着手"]], [14, 24, 14, 20, 18, 16, 22, 24, 16, 14, 12], 11),
        ("手順詳細", "1手順ごとの実施ステップと証跡を記録する。", ["手順ID", "Step", "操作・確認内容", "入力/コマンド", "期待結果", "異常時の対応", "証跡保存先", "注意事項"], [["RUN-001", 1, "バックアップジョブの実行結果を確認", "管理画面を開く", "成功ステータス", "担当リーダーへ連絡", "運用ログ", "個人情報を含む画面を共有しない"]], [14, 10, 32, 25, 24, 28, 24, 30], None),
        ("連絡先", "障害・判断・エスカレーション時の連絡先を管理する。", ["連絡先ID", "区分", "組織/担当", "氏名", "連絡手段", "対応時間帯", "一次/二次", "備考"], [["CNT-001", "障害", "運用チーム", "", "チャット/電話", "24x7", "一次", "重大度Aは電話"]], [14, 14, 22, 16, 22, 18, 14, 28], None),
    ],
)

save_workbook(
    "02_monitoring_design_template.xlsx",
    "監視設計書",
    "監視対象、メトリクス、閾値、通知、一次対応を定義する。",
    [
        ("監視項目", "監視対象と判定条件を定義する。", ["監視ID", "対象", "監視種別", "メトリクス/ログ", "収集間隔", "正常条件", "警告閾値", "異常閾値", "保持期間", "担当", "状態"], [["MON-001", "APIサーバー", "死活", "HTTP応答", "1分", "200応答", "3回連続失敗", "5回連続失敗", "90日", "運用担当", "未着手"]], [14, 22, 14, 24, 14, 24, 18, 18, 14, 16, 12], 11),
        ("アラート", "アラート発生時の通知先と対応を定義する。", ["アラートID", "監視ID", "重大度", "通知先", "通知方法", "一次対応", "エスカレーション条件", "抑止条件", "復旧確認", "備考"], [["ALT-001", "MON-001", "高", "運用当番", "メール/チャット", "サービス状態確認", "15分以内に復旧しない", "計画停止時間帯", "正常応答を確認", ""]], [14, 14, 12, 20, 18, 26, 28, 24, 24, 24], None),
        ("ダッシュボード", "運用状況を確認する画面・レポートを定義する。", ["画面ID", "画面名", "対象者", "表示指標", "更新頻度", "参照URL/場所", "アクセス権", "利用目的"], [["DASH-001", "サービス稼働状況", "運用責任者", "稼働率、エラー率", "リアルタイム", "", "運用管理者", "日次の状態確認"]], [14, 24, 20, 30, 16, 28, 20, 30], None),
    ],
)

save_workbook(
    "03_sla_slo_design_template.xlsx",
    "SLA/SLO定義書",
    "サービス品質の指標、目標、測定方法、報告方法を定義する。",
    [
        ("サービス定義", "対象サービスと提供時間、責任範囲を定義する。", ["サービスID", "サービス名", "対象範囲", "提供時間", "対象外時間", "利用者", "提供者", "前提・除外"], [["SVC-001", "業務申請サービス", "申請受付・照会", "平日8:00-20:00", "計画停止", "社内利用者", "運用チーム", "ネットワーク障害は対象外"]], [14, 24, 30, 20, 20, 20, 20, 32], None),
        ("SLI_SLO", "測定指標と目標値、未達時の対応を定義する。", ["目標ID", "サービスID", "SLI", "測定式", "測定期間", "目標値", "許容範囲", "データソース", "報告先", "未達時の対応", "状態"], [["SLO-001", "SVC-001", "可用性", "正常時間/提供時間", "月次", "99.9%", "99.5%未満で報告", "監視ログ", "サービス責任者", "原因分析と改善計画", "未着手"]], [14, 14, 20, 28, 14, 16, 20, 22, 20, 28, 12], 11),
        ("報告・責任", "測定結果の報告、責任者、見直し条件を定義する。", ["項目ID", "対象", "報告頻度", "報告様式", "作成者", "確認者", "承認者", "見直し条件", "証跡保存先"], [["SLA-REP-001", "SLO達成状況", "月次", "月次運用報告", "運用担当", "PM", "サービス責任者", "重大障害または目標変更", ""]], [16, 24, 16, 24, 18, 18, 18, 28, 24], None),
    ],
)

save_workbook(
    "04_incident_postmortem_template.xlsx",
    "インシデント報告書・ポストモーテム",
    "障害の影響、時系列、原因、対応、再発防止を記録する。",
    [
        ("概要", "インシデントの影響と基本情報を記録する。", ["インシデントID", "発生日", "検知日時", "復旧日時", "サービス", "重大度", "影響範囲", "影響時間", "報告者", "状態"], [["INC-001", "", "", "", "業務申請サービス", "高", "申請登録不可", "", "", "対応中"]], [16, 14, 14, 14, 24, 12, 32, 14, 16, 12], 10),
        ("時系列", "検知から復旧・連絡までの事実を時系列で記録する。", ["時刻", "実施者", "事象/判断", "対応内容", "結果", "証跡", "備考"], [["", "運用当番", "エラー率上昇を検知", "アラート確認と一次切り分け", "原因調査へ移行", "監視ログ", ""]], [16, 18, 30, 32, 24, 24, 28], None),
        ("原因_対策", "原因、影響、恒久対策、再発防止を記録する。", ["項目ID", "分析項目", "内容", "担当", "期限", "確認方法", "関連課題ID", "状態"], [["PM-001", "根本原因", "設定変更のレビュー不足", "開発リーダー", "", "レビュー記録確認", "", "対応中"], ["PM-002", "再発防止", "変更時の自動検証を追加", "開発リーダー", "", "CI結果", "", "未着手"]], [16, 18, 44, 20, 14, 24, 18, 12], 8),
    ],
)

save_workbook(
    "05_change_request_template.xlsx",
    "変更管理記録（RFC）",
    "変更要求の受付、影響分析、承認、実施、確認を一貫して管理する。",
    [
        ("変更要求", "変更の目的、範囲、影響、承認を記録する。", ["RFC ID", "申請日", "申請者", "変更概要", "変更理由", "対象", "希望日", "緊急度", "影響度", "承認者", "状態"], [["RFC-001", "", "", "監視閾値を変更", "誤検知削減", "監視設定", "", "通常", "中", "", "未着手"]], [14, 14, 16, 30, 28, 22, 14, 12, 12, 18, 12], 11),
        ("影響分析", "変更による影響と実施計画を確認する。", ["RFC ID", "影響領域", "影響内容", "リスク", "必要なテスト", "停止影響", "ロールバック方法", "担当", "確認者"], [["RFC-001", "運用", "アラート件数が減少", "異常見逃し", "閾値変更後の検知テスト", "なし", "変更前設定へ戻す", "", ""]], [14, 18, 32, 26, 30, 16, 28, 18, 18], None),
        ("実施結果", "実施、検証、承認、証跡を記録する。", ["RFC ID", "実施日時", "実施者", "実施結果", "検証結果", "問題・差戻し", "証跡", "承認/確認者", "状態"], [["RFC-001", "", "", "", "", "", "", "", "未着手"]], [14, 16, 18, 28, 28, 28, 28, 20, 12], 9),
    ],
)

save_workbook(
    "06_monthly_operations_report_template.xlsx",
    "月次運用報告書",
    "月次の稼働、障害、問い合わせ、変更、SLA/SLO、改善状況を報告する。",
    [
        ("サマリー", "月次の重要事項と総合判断を記録する。", ["報告月", "対象サービス", "総合評価", "重要な出来事", "経営・顧客への依頼", "作成者", "確認者", "承認者"], [["YYYY年MM月", "", "良/注意/要対応", "", "", "", "", ""]], [16, 24, 18, 42, 36, 18, 18, 18], None),
        ("運用実績", "稼働、障害、問い合わせ、変更の実績を記録する。", ["指標ID", "指標", "目標", "実績", "前月", "差異", "評価", "コメント"], [["OPS-001", "サービス稼働率", "99.9%", "", "", "", "", ""], ["OPS-002", "インシデント件数", "0件", "", "", "", "", ""], ["OPS-003", "変更件数", "", "", "", "", "", ""]], [16, 24, 16, 16, 16, 16, 16, 36], None),
        ("課題_改善", "継続課題、リスク、改善提案を記録する。", ["ID", "区分", "内容", "影響", "対応方針", "担当", "期限", "状態"], [["IMP-001", "改善", "定常作業を自動化", "作業時間削減", "自動化方式を検討", "", "", "未着手"]], [16, 14, 36, 24, 32, 18, 14, 12], 8),
    ],
)

save_workbook(
    "07_improvement_backlog_template.xlsx",
    "改善バックログ",
    "技術的負債、運用品質、プロセス、コスト、セキュリティの改善候補を継続管理する。",
    [
        ("改善項目", "改善の内容、効果、優先度、実施状況を管理する。", ["改善ID", "登録日", "区分", "改善内容", "背景・課題", "期待効果", "優先度", "工数見積", "担当", "期限", "状態"], [["IMP-001", "", "運用", "定常確認を自動化", "手作業による確認漏れ", "作業時間とミス削減", "高", "", "", "", "未着手"]], [16, 14, 14, 32, 32, 30, 12, 14, 18, 14, 12], 11),
        ("評価", "改善の実施結果と効果を評価する。", ["改善ID", "実施日", "実施内容", "実績効果", "期待との差異", "残課題", "評価者", "証跡", "状態"], [["IMP-001", "", "", "", "", "", "", "", "未着手"]], [16, 14, 32, 28, 24, 30, 18, 24, 12], 9),
    ],
)

save_workbook(
    "08_backup_dr_design_template.xlsx",
    "バックアップ・BCP/DR設計書",
    "バックアップ、復旧、代替運用、訓練、復旧目標を定義する。",
    [
        ("復旧目標", "業務・サービスごとの復旧目標を定義する。", ["対象ID", "対象サービス/データ", "重要度", "RTO", "RPO", "許容停止時間", "代替手段", "責任者", "承認者"], [["DR-001", "業務DB", "高", "4時間", "1時間", "4時間", "代替環境へ切替", "", ""]], [16, 26, 12, 14, 14, 18, 28, 18, 18], None),
        ("バックアップ", "バックアップ方式、頻度、保管、復元確認を定義する。", ["バックアップID", "対象", "方式", "頻度", "世代数", "保管場所", "暗号化", "復元テスト頻度", "確認者", "状態"], [["BKP-001", "業務DB", "フル+差分", "日次", "30世代", "別リージョン", "あり", "四半期", "運用担当", "未着手"]], [18, 22, 18, 14, 14, 24, 14, 22, 18, 12], 10),
        ("復旧手順_訓練", "障害時の復旧手順と訓練結果を記録する。", ["手順/訓練ID", "シナリオ", "開始条件", "手順概要", "成功条件", "実施日", "結果", "課題", "次回期限", "状態"], [["DR-TEST-001", "リージョン障害", "主要リージョン停止", "代替環境へ切替", "RTO/RPO達成", "", "", "", "", "未着手"]], [18, 24, 26, 32, 24, 14, 18, 30, 14, 12], 10),
    ],
)

save_workbook(
    "09_operations_handover_template.xlsx",
    "運用引き継ぎ・受入記録",
    "開発から運用への引き継ぎ対象、教育、受入条件、未完了事項を管理する。",
    [
        ("引き継ぎ項目", "運用開始に必要な資料、設定、作業、教育を一覧化する。", ["項目ID", "分類", "引き継ぎ対象", "内容", "提供元", "受領先", "期限", "確認方法", "証跡", "状態"], [["HDO-001", "資料", "Runbook", "定常運用・障害対応手順", "開発チーム", "運用チーム", "", "運用担当レビュー", "", "未着手"]], [16, 16, 24, 34, 18, 18, 14, 24, 24, 12], 10),
        ("教育_受入", "教育、演習、受入判定を記録する。", ["記録ID", "対象", "実施内容", "実施日", "参加者", "理解度/結果", "未解決事項", "受入者", "状態"], [["HDO-TRN-001", "運用担当", "障害一次対応の説明と演習", "", "", "", "", "", "未着手"]], [18, 22, 34, 14, 24, 24, 30, 18, 12], 9),
        ("未完了_連絡先", "残課題、問い合わせ先、エスカレーションを管理する。", ["ID", "区分", "内容", "影響", "対応者", "期限", "エスカレーション先", "状態"], [["HDO-OPEN-001", "残課題", "監視ダッシュボードの追加", "運用確認に影響", "", "", "サービス責任者", "未着手"]], [18, 14, 36, 24, 18, 14, 24, 12], 8),
    ],
)

print(f"生成完了: {len(list(OUTPUT_DIR.glob('*.xlsx')))} files")
