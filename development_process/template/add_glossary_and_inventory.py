from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BASE_DIR = Path(__file__).parent
NAVY = "1F4E78"
BLUE = "D9EAF7"
LIGHT = "F3F6F9"
WHITE = "FFFFFF"
GRAY = "666666"
THIN = Side(style="thin", color="B7C9D6")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

GLOSSARY = [
    ("Runbook", "定常運用、障害一次対応、復旧などを、担当者が同じ手順で実施できるように定義した手順書。", "運用手順、運用引き継ぎ"),
    ("SLA", "サービス提供者と利用者の間で合意したサービス水準、責任、対応条件。", "SLA/SLO定義、月次運用報告"),
    ("SLO", "サービス品質について設定する目標値。可用性、応答時間などの目標を定義する。", "SLA/SLO定義、監視設計"),
    ("SLI", "サービスの状態や品質を測定する具体的な指標。", "SLA/SLO定義、監視設計"),
    ("RTO", "Recovery Time Objective。障害発生からサービスを復旧するまでの目標時間。", "バックアップ・BCP/DR"),
    ("RPO", "Recovery Point Objective。障害時に許容するデータ損失時点の目標。", "バックアップ・BCP/DR"),
    ("BCP/DR", "事業継続計画と災害・大規模障害からの復旧計画。代替運用、復旧手順、訓練を含む。", "バックアップ・BCP/DR"),
    ("インシデント", "サービスの中断、品質低下、セキュリティ事象など、通常状態からの望ましくない逸脱。", "インシデント報告、監視設計"),
    ("ポストモーテム", "インシデントの事実、影響、原因、対応、再発防止を振り返る記録。個人の責任追及を目的としない。", "インシデント報告"),
    ("RFC", "Request for Change。変更要求を受付、影響分析、承認、実施、確認するための記録。", "変更管理"),
    ("変更管理", "変更要求の受付から影響分析、承認、実施、確認、記録までを管理する活動。", "変更管理、運用設計"),
    ("影響分析", "変更や障害が機能、データ、性能、セキュリティ、運用、納期などに与える影響を確認すること。", "変更管理、インシデント報告"),
    ("エスカレーション", "担当者やチームで解決できない課題や判断を、適切な責任者へ引き上げること。", "Runbook、監視、インシデント"),
    ("監視", "システムの状態、性能、ログ、業務結果などを継続的に観測すること。", "監視設計"),
    ("アラート", "監視で異常や閾値超過を検知し、通知や対応を開始するための信号。", "監視設計"),
    ("ロールバック", "問題発生時に、変更前のバージョンや設定へ戻すこと。", "変更管理、運用引き継ぎ"),
    ("運用引き継ぎ", "開発から運用へ、手順、監視、権限、連絡先、障害対応などを移管する活動。", "運用引き継ぎ"),
    ("改善バックログ", "技術的負債、運用品質、プロセス、コスト、セキュリティなどの改善候補を管理する一覧。", "改善バックログ、月次報告"),
    ("技術的負債", "短期的な判断や制約の結果として将来の変更・運用コストを増やす、未解消の技術上の課題。", "改善バックログ"),
    ("品質ゲート", "品質、リスク、未解決事項、承認状況を確認し、次工程へ進めるか判断する条件または場。", "運用受入、変更管理"),
    ("証跡", "実施、確認、判断、承認が行われたことを後から確認できる記録。", "全テンプレート"),
    ("テーラリング", "標準プロセスを案件の規模・リスクに応じて統合、簡略化、追加すること。", "全テンプレート"),
    ("重大度", "インシデントや不具合が利用者、業務、データなどに与える影響の大きさ。", "インシデント、変更管理"),
    ("優先度", "複数の課題や改善項目のうち、対応する順番や緊急性。重大度とは別に判断する。", "改善バックログ、月次報告"),
    ("改善効果", "改善を実施した結果として得られた、品質、時間、コスト、リスクなどの変化。", "改善バックログ、月次報告"),
]


def add_glossary_sheet(wb):
    if "用語集" in wb.sheetnames:
        del wb["用語集"]
    ws = wb.create_sheet("用語集")
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:C1")
    ws["A1"] = "運用・保守 用語集"
    ws["A1"].font = Font(bold=True, size=16, color=WHITE)
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A2"] = "このシートは、各ひな型で使用する運用・保守用語の意味をそろえるために使用します。"
    ws.merge_cells("A2:C2")
    ws["A2"].font = Font(italic=True, color=GRAY)
    ws["A2"].alignment = Alignment(wrap_text=True)
    headers = ["用語", "定義", "主な使用先"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(4, col, header)
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.border = BORDER
        cell.alignment = Alignment(wrap_text=True)
    for row, values in enumerate(GLOSSARY, 5):
        for col, value in enumerate(values, 1):
            cell = ws.cell(row, col, value)
            cell.border = BORDER
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=LIGHT)
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 80
    ws.column_dimensions["C"].width = 32
    ws.auto_filter.ref = f"A4:C{4 + len(GLOSSARY)}"
    ws.freeze_panes = "A5"


def add_inventory():
    records = [
        ("development_process/template/01_operation_runbook_template.xlsx", "運用手順、定常作業、障害一次対応、復旧手順の定義", "運用・保守テンプレート", "用語集あり"),
        ("development_process/template/02_monitoring_design_template.xlsx", "監視対象、メトリクス、閾値、アラート、一次対応の定義", "運用・保守テンプレート", "用語集あり"),
        ("development_process/template/03_sla_slo_design_template.xlsx", "サービス品質の指標、目標、測定、報告方法の定義", "運用・保守テンプレート", "用語集あり"),
        ("development_process/template/04_incident_postmortem_template.xlsx", "インシデントの影響、時系列、原因、再発防止の記録", "運用・保守テンプレート", "用語集あり"),
        ("development_process/template/05_change_request_template.xlsx", "変更要求、影響分析、承認、実施、検証の管理", "運用・保守テンプレート", "用語集あり"),
        ("development_process/template/06_monthly_operations_report_template.xlsx", "月次の稼働、障害、問い合わせ、変更、改善の報告", "運用・保守テンプレート", "用語集あり"),
        ("development_process/template/07_improvement_backlog_template.xlsx", "技術的負債、運用品質、プロセスなどの改善候補の管理", "運用・保守テンプレート", "用語集あり"),
        ("development_process/template/08_backup_dr_design_template.xlsx", "バックアップ、復旧目標、代替運用、訓練の定義", "運用・保守テンプレート", "用語集あり"),
        ("development_process/template/09_operations_handover_template.xlsx", "開発から運用への引き継ぎ、教育、受入、残課題の管理", "運用・保守テンプレート", "用語集あり"),
    ]
    wb = Workbook()
    ws = wb.active
    ws.title = "Excel一覧"
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:D1")
    ws["A1"] = "リポジトリ内 Excel ファイル一覧"
    ws["A1"].font = Font(bold=True, size=16, color=WHITE)
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws.merge_cells("A2:D2")
    ws["A2"] = "運用・保守テンプレート9件の用途を一覧化しています。用語集は各テンプレートと本一覧ブックに追加しています。"
    ws["A2"].alignment = Alignment(wrap_text=True)
    ws["A2"].font = Font(italic=True, color=GRAY)
    headers = ["ファイルパス", "用途", "区分", "用語集シート"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(4, col, header)
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.border = BORDER
    for row, values in enumerate(records, 5):
        for col, value in enumerate(values, 1):
            cell = ws.cell(row, col, value)
            cell.border = BORDER
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=LIGHT)
    widths = [58, 58, 24, 24]
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.auto_filter.ref = f"A4:D{4 + len(records)}"
    ws.freeze_panes = "A5"
    add_glossary_sheet(wb)
    wb.save(BASE_DIR / "00_excel_inventory_and_glossary.xlsx")


for path in sorted(BASE_DIR.glob("*.xlsx")):
    if path.name == "00_excel_inventory_and_glossary.xlsx":
        continue
    wb = load_workbook(path)
    add_glossary_sheet(wb)
    wb.save(path)

add_inventory()
print("用語集シート追加とExcel一覧作成が完了しました。")
